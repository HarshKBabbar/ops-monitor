"""Excel export of the NOTAM sheets (same colours as the dashboard)."""
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FILL = {
    "new": "FFFFFF", "pending": "FFF4C2", "done": "D7F0DC", "expiring": "F9D3D0", "archived": "FFFFFF",
}
LABEL = {
    "new": "Not reviewed", "pending": "Action pending", "done": "Done",
    "expiring": "Expiring", "archived": "Archived",
}
NEED = {"YES": "Yes", "NO": "No action needed", "EXT": "Extending"}


def fmt(ts):
    if not ts:
        return ""
    try:
        return datetime.fromisoformat(ts).strftime("%d %b %y %H:%MZ")
    except ValueError:
        return ts


def build(rows: list[dict], archive: bool) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Cancelled-Expired" if archive else "Active NOTAMs"
    head = ["Status", "NOTAM", "Location", "Aerodrome", "Valid From (UTC)", "Valid To (UTC)", "Schedule",
            "NOTAM Text", "Tags", "Perf. Critical", "Source", "Action Needed?", "Previous NOTAM", "Action Taken",
            "Updated By", "Updated At"]
    if archive:
        head += ["Archive Reason", "Archived On"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
        c.alignment = Alignment(vertical="center")
    for r in rows:
        vt = fmt(r["valid_to"]) or r["valid_to_text"]
        if r["valid_to"] and r["valid_to_text"] == "EST":
            vt += " EST"
        line = [LABEL[r["status"]], r["notam_id"], r["location"], r["airport_name"], fmt(r["valid_from"]), vt,
                r["schedule"], r["text"], ", ".join(r["tags"]), "Yes" if r["critical"] else "", r["sources"],
                NEED.get(r["action_needed"], ""), r["prev_notam"] or "", r["action_taken"], r["action_by"] or "",
                fmt(r["action_at"])]
        if archive:
            line += [r["archive_reason"], fmt(r["archived_at"])]
        ws.append(line)
        fill = PatternFill("solid", fgColor=FILL[r["status"]])
        for c in ws[ws.max_row]:
            c.fill = fill
            c.alignment = Alignment(wrap_text=True, vertical="top")
    widths = [13, 11, 9, 26, 16, 18, 16, 70, 18, 10, 10, 16, 14, 40, 14, 16, 34, 16]
    for i, w in enumerate(widths[: len(head)], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
